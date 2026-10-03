"""Synchronize the TrafficVision planning documents with the approved YOLO spec.

The script intentionally edits only the two named delivery documents.  It keeps
their existing Office layout and writes no model metrics, because no official
held-out-test result is available for the current project scope.
"""

from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ACTION_PLAN = DOCS / "TrafficVision_Project_Action_Plan.docx"
WBS = DOCS / "TrafficVision_Project_Work_Breakdown_Structure.xlsx"


def date(value: str) -> datetime:
    return datetime.strptime(value, "%Y-%m-%d")


def set_word_cell(cell, text: str) -> None:
    """Replace a table cell's content while retaining its paragraph/table style."""
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(12)


def update_action_plan() -> None:
    document = Document(ACTION_PLAN)
    document.paragraphs[0].text = "Kế hoạch dự án"
    for run in document.paragraphs[0].runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    overview = document.tables[0]
    details = {
        0: "Phát triển hệ thống thông minh",
        1: "RoadVision",
        2: "Phí Văn Nam /\nĐỗ Thị Vân Anh; Đỗ Hữu Nghị; Quản Văn Điệp",
        3: "TrafficVision – Nhận dạng biển báo giao thông Việt Nam bằng AI",
        5: (
            "Mục tiêu tổng quát. Xây dựng TrafficVision, ứng dụng web Streamlit chạy cục bộ "
            "cho phép phân tích ảnh và video biển báo giao thông bằng YOLO; ứng dụng được thiết kế "
            "để dùng baseline trước, sau đó thay thế an toàn bằng mô hình đã fine-tune trên dữ liệu biển báo Việt Nam.\n"
            "Mục tiêu 1 – Ứng dụng và suy luận. Hoàn thiện luồng tải ảnh/video, phát hiện và khoanh vùng, "
            "nhãn tiếng Việt theo manifest, độ tin cậy, thống kê, ảnh/video đã chú thích và CSV; suy luận "
            "ONNX CPU ổn định trên Windows/macOS.\n"
            "Mục tiêu 2 – Dữ liệu và huấn luyện. Kiểm định dữ liệu YOLO 82 lớp, EDA, snapshot bất biến, "
            "huấn luyện/fine-tune có checkpoint, đánh giá tập test độc lập, xuất ONNX và đóng gói Candidate.\n"
            "Mục tiêu 3 – An toàn vận hành và bàn giao. Bảo vệ Production bằng checksum, backup, promotion/rollback; "
            "hoàn thiện kiểm thử, README, Final Report và demo. Không đặt trước accuracy, mAP hay latency khi chưa đo."
        ),
        7: (
            "Phạm vi sử dụng. TrafficVision là đồ án minh họa quy trình phát triển hệ thống thông minh từ dữ liệu, "
            "huấn luyện và quản lý phiên bản mô hình đến suy luận web. Người dùng tải ảnh JPG/PNG/WebP hoặc video "
            "MP4/AVI/MOV, cấu hình confidence/IoU, xem kết quả và tải đầu ra đã chú thích cùng CSV.\n"
            "Hai giai đoạn. Giai đoạn 1 dùng YOLO11n baseline để kiểm chứng luồng end-to-end; baseline không được "
            "trình bày là mô hình biển báo Việt Nam. Giai đoạn 2 chỉ dùng dữ liệu đã qua Quality Gate để fine-tune, "
            "đánh giá và đưa Candidate lên Production sau thao tác review rõ ràng.\n"
            "Giới hạn. Không có webcam trực tiếp, tài khoản nhiều người dùng, cloud công cộng hay điều khiển phương tiện. "
            "Kết quả AI chỉ hỗ trợ tham khảo và có thể giảm chất lượng khi biển báo nhỏ, mờ, nghiêng hoặc che khuất."
        ),
        9: (
            "Bước 1 – Khởi tạo ứng dụng baseline. Phí Văn Nam thiết lập cấu trúc, cấu hình, model registry, manifest "
            "và ONNX baseline; Quản Văn Điệp triển khai giao diện Streamlit và các trạng thái hiển thị.\n"
            "Bước 2 – Tiếp nhận dữ liệu và Quality Gate. Đỗ Thị Vân Anh kiểm kê bộ dữ liệu biển báo Việt Nam; kiểm tra ảnh, "
            "nhãn YOLO, class ID, tọa độ, checksum và rò rỉ train/val/test; chỉ dữ liệu hợp lệ mới tạo snapshot/EDA.\n"
            "Bước 3 – Huấn luyện và Candidate. Đỗ Hữu Nghị cấu hình thiết bị, epoch, batch, checkpoint và resume an toàn; "
            "đánh giá độc lập, phân tích lỗi, xuất ONNX, kiểm tra parity và đóng gói Candidate.\n"
            "Bước 4 – Quản lý vòng đời model. Candidate chỉ được promotion khi manifest/checksum/evaluation hợp lệ; mọi promotion "
            "tạo backup và có rollback. Production không bị thay đổi bởi train hoặc fine-tune.\n"
            "Bước 5 – Hard-case correction. Nhóm tiếp nhận ảnh khó, gợi ý/sửa box và lớp, duyệt thủ công, tách challenge khỏi train/val "
            "và so sánh before/after trên cùng challenge.\n"
            "Bước 6 – Xác minh và bàn giao. Chạy unit/integration/UI tests, kiểm thử lỗi đầu vào, benchmark CPU, cập nhật README, "
            "Final Report, kịch bản demo và review cả nhóm."
        ),
        11: (
            "Nguồn và phạm vi. Bộ dữ liệu đích là dữ liệu biển báo giao thông Việt Nam theo định dạng YOLO, danh mục 82 lớp; "
            "nguồn, phiên bản và thời điểm tiếp nhận được lưu cùng artifact. YOLO11n COCO chỉ là baseline kỹ thuật, không dùng để "
            "tuyên bố độ chính xác biển báo Việt Nam.\n"
            "Kiểm định. Quality Gate chặn ảnh hỏng, dòng nhãn sai định dạng, class ID ngoài 0–81, tọa độ không hợp lệ và ảnh trùng "
            "giữa các split. Mất cân bằng lớp và box nhỏ là cảnh báo để review.\n"
            "Snapshot và truy vết. Snapshot chứa data.yaml, danh mục lớp, checksum, seed, báo cáo validation/EDA và provenance; "
            "tập test độc lập không bị gộp vào train/val. Hard-case challenge_only chỉ phục vụ so sánh, không dùng huấn luyện."
        ),
        13: (
            "Gói ứng dụng. Ứng dụng Streamlit có các trang Phân tích, Lịch sử, Thống kê, Huấn luyện AI, Thông tin model và Thiết lập; "
            "xử lý ảnh/video theo giới hạn bộ nhớ, lưu kết quả đã chú thích và CSV.\n"
            "Gói dữ liệu/model. Báo cáo Quality Gate, EDA, snapshot, cấu hình training, log/checkpoint, evaluation, ONNX, manifest, "
            "checksum, Candidate và backup/rollback history.\n"
            "Gói cải tiến. Correction store, hàng chờ duyệt, collection goal, snapshot fine-tune và challenge report before/after; "
            "không tự động promotion Candidate.\n"
            "Hồ sơ bàn giao. Source code, test evidence, README, WBS, Action Plan, Final Report, slide/kịch bản demo và known limitations."
        ),
        15: (
            "Phí Văn Nam – Trưởng nhóm, kiến trúc và tích hợp. Quản lý WBS/gate; phụ trách cấu trúc hệ thống, model registry, "
            "integration, promotion/rollback, tổng hợp báo cáo và rehearsal demo.\n"
            "Đỗ Thị Vân Anh – Dữ liệu. Phụ trách tiếp nhận dữ liệu, kiểm kê, Quality Gate, repair có kiểm soát, EDA, snapshot, "
            "catalog và provenance.\n"
            "Đỗ Hữu Nghị – Huấn luyện và đánh giá. Phụ trách cấu hình training, checkpoint/resume, fine-tune, đánh giá test độc lập, "
            "phân tích lỗi, ONNX/parity và Candidate.\n"
            "Quản Văn Điệp – Giao diện và QA. Phụ trách giao diện Streamlit, luồng ảnh/video, trạng thái lỗi, accessibility, browser/UI tests, "
            "hướng dẫn sử dụng và hỗ trợ demo.\n"
            "Trách nhiệm chung. Cả nhóm review dữ liệu, Candidate, kịch bản demo và Final Report; không tự thay đổi split, class mapping "
            "hay Production ngoài quy trình đã chốt."
        ),
    }
    for row, text in details.items():
        set_word_cell(overview.cell(row, 1), text)

    schedule = (
        "Tuần 1 — 28/09–04/10: Ứng dụng baseline và nền tảng. Nam/Điệp hoàn thiện registry, ONNX baseline, phân tích ảnh/video, "
        "đầu ra CSV và các trang Streamlit; Vân Anh bắt đầu kiểm kê dữ liệu. Gate: demo end-to-end trên CPU và cảnh báo baseline rõ ràng.\n"
        "Tuần 2 — 05–11/10: Dữ liệu và Quality Gate. Vân Anh kiểm định, repair có kiểm soát, EDA và snapshot; Điệp/Nam hoàn thiện "
        "UI huấn luyện, model info và thiết lập. Gate: snapshot không có lỗi blocking, sẵn sàng huấn luyện.\n"
        "Tuần 3 — 12–18/10: Huấn luyện, Candidate và hard cases. Nghị/Nam chạy training hoặc resume an toàn, đánh giá độc lập, xuất ONNX, "
        "parity và Candidate; nhóm mở correction queue, review và challenge. Gate: Candidate có bằng chứng đầy đủ để review, Production không đổi.\n"
        "Tuần 4 — 19–25/10: Xác minh và bàn giao. Hoàn thiện challenge report, test tự động/UI, benchmark CPU, README, Final Report, slide, "
        "demo và group review. Gate: bằng chứng truy vết đầy đủ, demo local end-to-end. Bàn giao dự kiến 26/10."
    )
    set_word_cell(document.tables[1].cell(1, 1), schedule)
    document.save(ACTION_PLAN)


def copy_row_style(sheet, source_row: int, target_row: int) -> None:
    for column in range(1, 31):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def update_wbs() -> None:
    workbook = load_workbook(WBS)
    sheet = workbook["WBS"]

    sheet["E7"] = date("2026-10-01")
    sheet["C9"] = "TrafficVision – Nhận dạng biển báo giao thông Việt Nam bằng YOLO"

    project_days = [
        "2026-09-28", "2026-09-29", "2026-09-30", "2026-10-01", "2026-10-04",
        "2026-10-05", "2026-10-06", "2026-10-07", "2026-10-08", "2026-10-11",
        "2026-10-12", "2026-10-13", "2026-10-14", "2026-10-15", "2026-10-18",
        "2026-10-19", "2026-10-20", "2026-10-21", "2026-10-22", "2026-10-25",
    ]
    for column, value in enumerate(project_days, start=11):
        sheet.cell(13, column).value = date(value)
        sheet.cell(13, column).number_format = "dd/mm"

    entries = [
        ("US-01 – Ứng dụng baseline và suy luận", "2026-09-28", "2026-10-04", "Phí Văn Nam + Quản Văn Điệp", "Baseline ONNX và demo ảnh/video", 1, 5),
        ("US-01.01 – Registry, manifest và ONNX baseline", "2026-09-28", "2026-09-29", "Phí Văn Nam", "Model registry; manifest; checksum", 1, 2),
        ("US-01.02 – Phân tích ảnh/video và đầu ra", "2026-09-30", "2026-10-04", "Quản Văn Điệp + Phí Văn Nam", "Ảnh/video chú thích; CSV; lịch sử", 3, 5),
        ("US-01.03 – Cảnh báo baseline và kiểm thử luồng", "2026-10-01", "2026-10-04", "Quản Văn Điệp", "UI states; cảnh báo; smoke test", 4, 5),
        ("US-02 – Dữ liệu biển báo Việt Nam và Quality Gate", "2026-09-28", "2026-10-08", "Đỗ Thị Vân Anh + Đỗ Hữu Nghị", "Dataset hợp lệ và snapshot bất biến", 1, 9),
        ("US-02.01 – Tiếp nhận, kiểm kê và catalog 82 lớp", "2026-09-28", "2026-09-30", "Đỗ Thị Vân Anh", "Nguồn; cấu trúc; mapping 82 lớp", 1, 3),
        ("US-02.02 – Validate, repair và chống leakage", "2026-10-04", "2026-10-06", "Đỗ Thị Vân Anh", "Quality Gate; repair manifest; quarantine", 5, 7),
        ("US-02.03 – EDA và tạo snapshot", "2026-10-07", "2026-10-08", "Đỗ Thị Vân Anh", "EDA report; data.yaml; snapshot manifest", 8, 9),
        ("US-03 – Streamlit AI Experiment Lab và vận hành", "2026-10-05", "2026-10-18", "Quản Văn Điệp + Phí Văn Nam", "UI huấn luyện và quản lý model", 6, 15),
        ("US-03.01 – Trang Phân tích, Lịch sử và Thống kê", "2026-10-05", "2026-10-06", "Quản Văn Điệp", "Các trang Streamlit và trạng thái lỗi", 6, 7),
        ("US-03.02 – UI Quality Gate, EDA và training", "2026-10-07", "2026-10-13", "Quản Văn Điệp + Đỗ Hữu Nghị", "4 bước Data–Gate–Train–Evaluate", 8, 12),
        ("US-03.03 – Model info, settings, backup/rollback", "2026-10-12", "2026-10-18", "Phí Văn Nam + Quản Văn Điệp", "Manifest; settings; lifecycle controls", 11, 15),
        ("US-04 – Huấn luyện, đánh giá và Candidate", "2026-09-28", "2026-10-18", "Đỗ Hữu Nghị + Phí Văn Nam", "Run có checkpoint và Candidate reviewable", 1, 15),
        ("US-04.01 – Runtime phần cứng và cấu hình an toàn", "2026-09-28", "2026-10-04", "Đỗ Hữu Nghị", "Device/batch/workers; CPU fallback", 1, 5),
        ("US-04.02 – Train/fine-tune, checkpoint và resume", "2026-10-05", "2026-10-15", "Đỗ Hữu Nghị", "Run log; best.pt; last.pt; resume", 6, 14),
        ("US-04.03 – Test độc lập, ONNX và Candidate", "2026-10-14", "2026-10-18", "Đỗ Hữu Nghị + Phí Văn Nam", "Metrics thực tế; parity; candidate package", 13, 15),
        ("US-05 – Hard-case correction và fine-tune có kiểm soát", "2026-10-12", "2026-10-21", "Phí Văn Nam + Đỗ Hữu Nghị + Quản Văn Điệp", "Correction queue và challenge report", 11, 18),
        ("US-05.01 – Upload, gợi ý/sửa box và lớp", "2026-10-12", "2026-10-14", "Quản Văn Điệp", "Correction item; annotation; preview", 11, 13),
        ("US-05.02 – Duyệt, collection goal và snapshot", "2026-10-15", "2026-10-19", "Phí Văn Nam + Đỗ Thị Vân Anh", "Approval queue; goal; provenance snapshot", 14, 16),
        ("US-05.03 – Fine-tune mới và challenge before/after", "2026-10-20", "2026-10-21", "Đỗ Hữu Nghị", "Candidate mới; challenge report", 17, 18),
        ("US-06 – Xác minh, báo cáo và bàn giao", "2026-10-19", "2026-10-25", "Cả nhóm", "Evidence, Final Report và demo", 16, 20),
        ("US-06.01 – Test tự động, UI và lỗi biên", "2026-10-19", "2026-10-20", "Quản Văn Điệp + Phí Văn Nam", "Pytest/Ruff; UI evidence; failure paths", 16, 17),
        ("US-06.02 – Benchmark CPU, README và limits", "2026-10-21", "2026-10-22", "Phí Văn Nam + Đỗ Hữu Nghị", "Benchmark; hướng dẫn; known limitations", 18, 19),
        ("US-06.03 – Final Report, slide, demo và review", "2026-10-22", "2026-10-25", "Cả nhóm", "Báo cáo; slide; kịch bản và review cuối", 19, 20),
    ]

    for offset, entry in enumerate(entries):
        row = 14 + offset
        style_source = 14 + (offset % 4)
        copy_row_style(sheet, style_source, row)
        task, start, end, owner, deliverable, day_start, day_end = entry
        for column in range(2, 31):
            sheet.cell(row, column).value = None
        level = offset % 4
        sheet.cell(row, 2 + level).value = task
        sheet.cell(row, 6).value = date(start)
        sheet.cell(row, 7).value = date(end)
        sheet.cell(row, 8).value = owner
        sheet.cell(row, 9).value = deliverable
        sheet.cell(row, 10).value = "In-progress"
        for column in (6, 7):
            sheet.cell(row, column).number_format = "dd/mm/yyyy"
        for column in range(11, 31):
            if day_start <= column - 10 <= day_end:
                sheet.cell(row, column).value = "X"

    for row in range(38, 46):
        for column in range(2, 31):
            sheet.cell(row, column).value = None

    workbook.save(WBS)


def main() -> None:
    update_action_plan()
    update_wbs()
    print("ACTION_PLAN_YOLO_SCOPE=PASS")
    print("WBS_YOLO_SCOPE=PASS")


if __name__ == "__main__":
    main()
