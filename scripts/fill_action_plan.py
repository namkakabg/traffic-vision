import sys
import os
import docx
from docx.shared import Pt, RGBColor

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_text(cell, paragraphs_data, font_name="SamsungOne 400", font_size=Pt(10.5), default_color=RGBColor(0x22, 0x22, 0x22)):
    """
    paragraphs_data: list of paragraphs.
    Each paragraph is a list of runs: (text, is_bold, optional_color)
    or a simple string.
    """
    # Remove extra paragraphs in cell except first
    p_first = cell.paragraphs[0]
    p_first.text = ""
    p_first.paragraph_format.line_spacing = 1.15
    p_first.paragraph_format.space_after = Pt(4)
    
    # Remove any existing additional paragraphs
    for p_extra in cell.paragraphs[1:]:
        p_extra._element.getparent().remove(p_extra._element)
        
    for p_idx, p_data in enumerate(paragraphs_data):
        if p_idx == 0:
            p = p_first
        else:
            p = cell.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)
            
        if isinstance(p_data, str):
            p_data = [(p_data, False, default_color)]
            
        for run_info in p_data:
            if isinstance(run_info, str):
                text = run_info
                bold = False
                color = default_color
            else:
                text = run_info[0]
                bold = run_info[1] if len(run_info) > 1 else False
                color = run_info[2] if len(run_info) > 2 else default_color
                
            run = p.add_run(text)
            run.font.name = font_name
            run.font.size = font_size
            run.font.bold = bold
            run.font.color.rgb = color

def fill_action_plan(src_path, dst_path):
    doc = docx.Document(src_path)
    
    t0 = doc.tables[0]
    # Row 1: Tên nhóm
    set_cell_text(t0.rows[1].cells[1], [
        [("TrafficVision", True)]
    ])
    
    # Row 2: Trưởng nhóm / Các thành viên nhóm
    set_cell_text(t0.rows[2].cells[1], [
        [("Phí Văn Nam (Trưởng nhóm)", True), ("\nĐỗ Thị Vân Anh, Đỗ Hữu Nghị, Quản Văn Điệp", False)]
    ])
    
    # Row 3: Tên đề tài
    set_cell_text(t0.rows[3].cells[1], [
        [("TrafficVision – Hệ thống nhận dạng biển báo giao thông Việt Nam bằng AI tối ưu suy luận trên CPU", True)]
    ])
    
    # Row 5: Mục tiêu (merged)
    muc_tieu = [
        [("1. Nhận dạng chính xác 82 lớp biển báo giao thông Việt Nam: ", True),
         ("Xây dựng mô hình thị giác máy tính nhận dạng và định vị chính xác 82 lớp biển báo giao thông đường bộ theo Quy chuẩn kỹ thuật quốc gia QCVN 41:2019/BGTVT trên cả ảnh tĩnh và luồng video.", False)],
        [("2. Tối ưu hóa suy luận trên phần cứng CPU phổ thông: ", True),
         ("Đóng gói mô hình dưới dạng ONNX Runtime đồ thị tĩnh 640x640 px, khai thác tập chỉ thị AVX2 đa luồng, đảm bảo tốc độ đáp ứng gần thời gian thực (~15 FPS cho bản Nano) mà không đòi hỏi GPU đắt tiền.", False)],
        [("3. Xây dựng quy trình MLOps khép kín và tin cậy: ", True),
          ("Thiết lập cổng Quality Gate 5 tiêu chí chặn rác dữ liệu, module Dataset Repair sửa lỗi bounding box, snapshot dữ liệu bất biến, tiến trình huấn luyện ngầm độc lập có Safe Resume, kiểm định sai lệch PyTorch vs ONNX, đóng gói Candidate và thăng cấp an toàn (Safe Atomic Swap) kèm cơ chế Rollback 1-click.", False)],
        [("4. Cung cấp sản phẩm phần mềm hoàn chỉnh: ", True),
         ("Xây dựng Web Dashboard Streamlit trực quan, hỗ trợ kéo thả ảnh/video, hiển thị bounding box tiếng Việt chuẩn hóa, thanh tiến trình thời gian thực (@st.fragment), xuất tệp CSV và lưu lịch sử CSDL SQLite.", False)]
    ]
    set_cell_text(t0.rows[5].cells[0], muc_tieu)
    if t0.rows[5].cells[1] != t0.rows[5].cells[0]:
        set_cell_text(t0.rows[5].cells[1], muc_tieu)
        
    # Row 7: Tóm tắt
    tom_tat = [
        [("TrafficVision ", True),
         ("là hệ thống thị giác máy tính thông minh giải quyết bài toán tự động nhận diện biển báo giao thông đường bộ tại Việt Nam. Dự án triển khai mô hình học sâu tiên tiến YOLO11 (Ultralytics) gồm hai cấu hình: ", False),
         ("YOLO11m (Medium) ", True),
         ("cho độ chính xác cao và ", False),
         ("YOLO11n (Nano) ", True),
         ("siêu nhẹ tối ưu độ trễ cho biên. Bộ dữ liệu chuẩn hóa gồm 10.129 ảnh và 19.700 bounding boxes bao phủ 82 lớp biển báo, được kiểm định qua Quality Gate 5 tiêu chí và đóng gói Snapshot bất biến.", False)],
        [("Mô hình được xuất sang định dạng ONNX Runtime đồ thị tĩnh và suy luận tối ưu trên CPU. Kết quả thực nghiệm trên tập kiểm thử độc lập (held-out test split) gồm 1.016 ảnh cho thấy mô hình YOLO11m đạt kết quả vượt bậc: ", False),
         ("mAP50 = 98.03%, Precision = 96.16%, Recall = 96.28%, F1-score = 96.22%", True),
         (", vượt qua kiểm tra sai số số học PyTorch vs ONNX (max_abs_diff = 0.000854 <= 1e-3) và đã chính thức thăng cấp lên Production. Bản YOLO11n đạt mAP50 = 92.47% với độ trễ CPU chỉ 68.26 ms/ảnh (~14.65 FPS). Hệ thống được bảo chứng độ tin cậy bởi bộ kiểm thử tự động 215 tests, giao diện Web Streamlit đa chức năng và đầy đủ tài liệu kỹ thuật.", False)]
    ]
    set_cell_text(t0.rows[7].cells[0], tom_tat)
    if t0.rows[7].cells[1] != t0.rows[7].cells[0]:
        set_cell_text(t0.rows[7].cells[1], tom_tat)
        
    # Row 9: Phương pháp
    phuong_phap = [
        [("• Kiến trúc mô hình: ", True),
         ("Sử dụng kiến trúc YOLO11 tiên tiến với cơ chế Attention không gian C2PSA và hàm mất mát phân bố tiêu điểm DFL Loss, giúp phát hiện hiệu quả các biển báo nhỏ ở xa và trong điều kiện ánh sáng phức tạp.", False)],
        [("• Quản lý chất lượng dữ liệu: ", True),
         ("Áp dụng cổng Quality Gate 5 tiêu chí (Corrupt Image, Format Line, Class ID 0-81, BBox Coords [0, 1], Data Leakage SHA-256) kết hợp module Dataset Repair tự động clamping tọa độ và loại bỏ box rỗng trước khi đóng gói Snapshot bất biến.", False)],
        [("• Huấn luyện ngầm & Khả năng chịu lỗi: ", True),
         ("Huấn luyện qua tiến trình nền độc lập (background runner) với cấu hình workers=0 tối ưu trên Windows, hỗ trợ Checkpoint & Safe Resume phục hồi an toàn từ last.pt.", False)],
        [("• Tối ưu suy luận & Parity Gate: ", True),
         ("Xuất mô hình sang ONNX đồ thị tĩnh 640x640, suy luận qua ONNX Runtime CPU với tập lệnh AVX2 đa luồng, kiểm định sai số số học cực đại PyTorch vs ONNX (max_abs_diff <= 1e-3).", False)],
        [("• MLOps Model Registry: ", True),
         ("Quản lý phiên bản mô hình qua manifest.json động và mã băm SHA-256; cơ chế Safe Atomic Swap đảm bảo cập nhật mô hình production không gián đoạn dịch vụ cùng khả năng Rollback 1-click tức thì.", False)]
    ]
    set_cell_text(t0.rows[9].cells[0], phuong_phap)
    if t0.rows[9].cells[1] != t0.rows[9].cells[0]:
        set_cell_text(t0.rows[9].cells[1], phuong_phap)

    t1 = doc.tables[1]
    # Row 1: Dữ liệu
    du_lieu = [
        [("• Nguồn dữ liệu: ", True),
         ("Bộ dữ liệu star092304/Traffic-sign-detection-VietNam từ Hugging Face kết hợp ảnh chụp thực tế camera hành trình tại Việt Nam, chuẩn hóa theo danh mục 82 lớp biển báo của Quy chuẩn QCVN 41:2019/BGTVT.", False)],
        [("• Quy mô dữ liệu: ", True),
         ("Tổng cộng 10.129 ảnh với 19.700 bounding boxes, phân bổ thành: Tập Train (8.098 ảnh / 15.671 boxes), Tập Val (1.015 ảnh / 2.059 boxes) và Tập Test độc lập (1.016 ảnh / 1.970 boxes).", False)],
        [("• Tiền xử lý & Snapshot: ", True),
         ("Dữ liệu thô đưa vào staging, vượt qua cổng kiểm định 5 tiêu chí, sửa lỗi tọa độ qua Dataset Repair và đóng gói vào Snapshot bất biến snapshot_20260930_165016 kèm tệp data.yaml và báo cáo EDA toàn diện (phân bố nhãn, kích thước bbox COCO small/medium/large, tỉ lệ khung hình).", False)]
    ]
    set_cell_text(t1.rows[1].cells[0], du_lieu)
    if t1.rows[1].cells[1] != t1.rows[1].cells[0]:
        set_cell_text(t1.rows[1].cells[1], du_lieu)
        
    # Row 3: Đầu ra mong muốn đạt được
    dau_ra = [
        [("• Kết quả mô hình AI: ", True),
         ("Mô hình Production (YOLO11m) nhận diện 82 lớp biển báo đạt mAP50 = 98.03%, F1-score = 96.22%, sai số số học PyTorch vs ONNX = 0.000854 <= 1e-3. Mô hình Candidate (YOLO11n) đạt mAP50 = 92.47% với độ trễ CPU chỉ 68.26 ms/ảnh (~14.65 FPS).", False)],
        [("• Gói sản phẩm phần mềm: ", True),
         ("Ứng dụng Web Streamlit hoàn chỉnh cho phép phân tích ảnh và video trực quan, hiển thị nhãn tiếng Việt chuẩn hóa, thanh tiến trình tự động làm mới, xuất tệp CSV thống kê và lưu trữ lịch sử phiên vào CSDL SQLite.", False)],
        [("• Artifact MLOps & Chất lượng: ", True),
         ("Toàn bộ artifact bất biến (Snapshot, Checkpoint, ONNX, Manifest SHA-256, test_metrics.json, benchmark.json); bộ kiểm thử tự động đạt 215/215 tests (213 passed, 2 skipped).", False)],
        [("• Lợi ích thực tiễn: ", True),
         ("Hỗ trợ đắc lực cho các hệ thống cảnh báo lái xe an toàn (ADAS); tự động hóa công tác khảo sát, số hóa và duy tu biển báo của cơ quan quản lý giao thông đô thị với chi phí phần cứng tối thiểu.", False)]
    ]
    set_cell_text(t1.rows[3].cells[0], dau_ra)
    if t1.rows[3].cells[1] != t1.rows[3].cells[0]:
        set_cell_text(t1.rows[3].cells[1], dau_ra)
        
    # Row 5: Vai trò của từng thành viên
    vai_tro = [
        [("• Phí Văn Nam (Trưởng nhóm): ", True),
         ("Thiết kế kiến trúc hệ thống tổng thể, MLOps pipeline, Model Registry, cơ chế Safe Atomic Swap & Rollback, tích hợp CSDL SQLite (history.db), điều phối tiến độ và tổng hợp báo cáo dự án.", False)],
        [("• Đỗ Thị Vân Anh (Kỹ sư Dữ liệu): ", True),
         ("Thu thập, tiền xử lý và chuẩn hóa dữ liệu 82 lớp biển báo theo QCVN 41:2019; xây dựng cổng Quality Gate 5 tiêu chí chặn, module Dataset Repair, phân tích EDA và đóng gói Snapshot bất biến.", False)],
        [("• Đỗ Hữu Nghị (Kỹ sư AI/ML): ", True),
         ("Cấu hình huấn luyện YOLO11 (Nano và Medium), tối ưu hóa suy luận ONNX trên CPU, kiểm tra sai số số học PyTorch vs ONNX, benchmark độ trễ/FPS phần cứng và đóng gói Model Candidate.", False)],
        [("• Quản Văn Điệp (Kỹ sư Fullstack / QA): ", True),
         ("Phát triển Web Dashboard Streamlit đa chức năng, xử lý luồng đa phương tiện ảnh/video, cơ chế cập nhật tiến trình thời gian thực (@st.fragment), xây dựng bộ kiểm thử tự động 215 tests và hoàn thiện tài liệu hướng dẫn.", False)],
        [("• Trách nhiệm chung: ", True),
         ("Thẩm định chéo dữ liệu, đánh giá kết quả thực nghiệm, tham gia diễn tập bảo vệ đồ án và cam kết tính trung thực 100% của mọi số liệu trích xuất từ artifact.", False)]
    ]
    set_cell_text(t1.rows[5].cells[0], vai_tro)
    if t1.rows[5].cells[1] != t1.rows[5].cells[0]:
        set_cell_text(t1.rows[5].cells[1], vai_tro)

    t2 = doc.tables[2]
    # Row 1: Tóm tắt lịch trình
    lich_trinh = [
        [("• Tuần 1 (28/09 – 02/10/2026) – Khởi tạo Baseline & Nền tảng MLOps: ", True),
         ("Xây dựng khung ứng dụng Streamlit end-to-end; tích hợp mô hình YOLO11n ONNX baseline trên CPU; thiết kế Model Registry, SQLite history; hoàn thiện catalog 82 lớp biển báo Việt Nam, cổng Quality Gate 5 tiêu chí, module Dataset Repair, đóng gói snapshot 10.129 ảnh và xuất báo cáo EDA. (Hoàn thành 100%)", False)],
        [("• Tuần 2 (05/10 – 09/10/2026) – Huấn luyện chuyên sâu & Tối ưu hóa: ", True),
         ("Hoàn thiện Training Engine ngầm với tính năng Safe Resume; huấn luyện 50 epochs cho cả hai cấu hình YOLO11n và YOLO11m; xuất mô hình ONNX đồ thị tĩnh [1, 3, 640, 640]; kiểm định sai số số học PyTorch vs ONNX (max_abs_diff <= 1e-3); đo đạc Benchmark CPU FPS và Latency; đóng gói Candidate hoàn chỉnh. (Hoàn thành 100%)", False)],
        [("• Tuần 3 (12/10 – 16/10/2026) – Đánh giá Độc lập, Thăng cấp & UAT: ", True),
         ("Đánh giá mô hình trên tập kiểm thử độc lập 1.016 ảnh (YOLO11m đạt mAP50 = 98.03%, F1 = 96.22%); thực hiện Safe Atomic Swap thăng cấp YOLO11m lên Production; kiểm thử cơ chế sao lưu tự động và Rollback 1-click; kiểm thử chấp nhận người dùng (UAT) trên ảnh thực tế và video camera hành trình. (Hoàn thành 100%)", False)],
        [("• Tuần 4 (19/10 – 23/10/2026) – Kiểm thử toàn diện & Bàn giao: ", True),
         ("Chạy bộ kiểm thử tự động hoàn chỉnh 215 tests (213 passed, 2 skipped); hoàn thiện hồ sơ bàn giao gồm Sổ tay Hướng dẫn sử dụng, Slide thuyết trình bảo vệ (14 slides), Cẩm nang Q&A theo 4 thành viên, Tài liệu kỹ thuật huấn luyện chuyên sâu và Báo cáo dự án cuối khóa. (Hoàn thành 100%)", False)],
        [("• Giai đoạn mở rộng sau bàn giao: ", True),
         ("Nghiên cứu tích hợp luồng RTSP camera hành trình trực tiếp thời gian thực, phát triển module cảnh báo giọng nói tiếng Việt (Text-to-Speech) và gắn tọa độ vị trí biển báo lên bản đồ số GPS.", False)]
    ]
    set_cell_text(t2.rows[1].cells[0], lich_trinh)
    if t2.rows[1].cells[1] != t2.rows[1].cells[0]:
        set_cell_text(t2.rows[1].cells[1], lich_trinh)
        
    # Row 3: Nhận xét của giáo viên
    nhan_xet = [
        [("(Dành cho giảng viên đánh giá và ghi nhận xét về kế hoạch hành động của nhóm)", False, RGBColor(0x88, 0x88, 0x88))]
    ]
    set_cell_text(t2.rows[3].cells[0], nhan_xet)
    if t2.rows[3].cells[1] != t2.rows[3].cells[0]:
        set_cell_text(t2.rows[3].cells[1], nhan_xet)

    doc.save(dst_path)
    print(f"Successfully saved Action Plan to: {dst_path}")

if __name__ == "__main__":
    fill_action_plan("docs/template/backup_original/Project_Action Plan.docx", "docs/template/Project_Action Plan.docx")
